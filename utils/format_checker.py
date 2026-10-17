# -*- coding: utf-8 -*-
"""
论文/作业格式检查与修复核心模块
支持多种格式模板、中英文字体分别设置、可检查可自动修复、记录修改详情
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re


# ============================================================
# 常用字体列表（供前端下拉选择）
# ============================================================
COMMON_CN_FONTS = [
    "宋体", "黑体", "楷体", "仿宋", "微软雅黑",
    "隶书", "幼圆", "华文行楷", "华文新魏", "华文中宋",
    "方正姚体", "方正舒体", "STSong", "STKaiti", "STFangsong"
]

COMMON_EN_FONTS = [
    "Times New Roman", "Arial", "Calibri", "Cambria",
    "Georgia", "Verdana", "Courier New", "Tahoma",
    "Garamond", "Book Antiqua", "Century", "Palatino Linotype"
]

# 字号对照表（中文字号 -> pt）
FONT_SIZE_MAP = {
    "初号": 42, "小初": 36,
    "一号": 26, "小一": 24,
    "二号": 22, "小二": 18,
    "三号": 16, "小三": 15,
    "四号": 14, "小四": 12,
    "五号": 10.5, "小五": 9,
    "六号": 7.5, "小六": 6.5,
}
FONT_SIZE_OPTIONS = [(name, pt) for name, pt in FONT_SIZE_MAP.items()]


def get_font_options():
    """返回前端用的字体选项"""
    return {
        "cn_fonts": COMMON_CN_FONTS,
        "en_fonts": COMMON_EN_FONTS,
        "sizes": [{"name": n, "pt": p} for n, p in FONT_SIZE_MAP.items()],
    }


# ============================================================
# 格式模板定义
# ============================================================
FORMAT_TEMPLATES = {
    "general": {
        "name": "通用学术论文",
        "description": "适用于一般课程论文、学术报告",
        "paper_width": 21.0, "paper_height": 29.7,
        "top_margin": 2.54, "bottom_margin": 2.54,
        "left_margin": 3.17, "right_margin": 3.17,
        "body_font": "宋体", "body_en_font": "Times New Roman",
        "body_size": 12, "body_line_spacing": 1.5,
        "body_first_line_indent": 24, "body_alignment": "justify",
        "h1_font": "黑体", "h1_en_font": "Times New Roman",
        "h1_size": 16, "h1_bold": False, "h1_alignment": "center",
        "h1_space_before": 12, "h1_space_after": 12,
        "h2_font": "黑体", "h2_en_font": "Times New Roman",
        "h2_size": 14, "h2_bold": False, "h2_alignment": "left",
        "h2_space_before": 6, "h2_space_after": 6,
        "h3_font": "黑体", "h3_en_font": "Times New Roman",
        "h3_size": 12, "h3_bold": True, "h3_alignment": "left",
        "ref_font": "宋体", "ref_en_font": "Times New Roman",
        "ref_size": 10.5, "ref_line_spacing": 1.25,
        "keyword_font": "黑体", "keyword_en_font": "Times New Roman",
        "keyword_size": 12,
    },
    "thesis": {
        "name": "本科毕业论文",
        "description": "适用于本科毕业论文、毕业设计",
        "paper_width": 21.0, "paper_height": 29.7,
        "top_margin": 2.5, "bottom_margin": 2.5,
        "left_margin": 3.0, "right_margin": 2.5,
        "body_font": "宋体", "body_en_font": "Times New Roman",
        "body_size": 12, "body_line_spacing": 1.5,
        "body_first_line_indent": 24, "body_alignment": "justify",
        "h1_font": "黑体", "h1_en_font": "Times New Roman",
        "h1_size": 18, "h1_bold": False, "h1_alignment": "center",
        "h1_space_before": 18, "h1_space_after": 18,
        "h2_font": "黑体", "h2_en_font": "Times New Roman",
        "h2_size": 14, "h2_bold": False, "h2_alignment": "left",
        "h2_space_before": 12, "h2_space_after": 6,
        "h3_font": "黑体", "h3_en_font": "Times New Roman",
        "h3_size": 12, "h3_bold": True, "h3_alignment": "left",
        "ref_font": "宋体", "ref_en_font": "Times New Roman",
        "ref_size": 10.5, "ref_line_spacing": 1.25,
        "keyword_font": "黑体", "keyword_en_font": "Times New Roman",
        "keyword_size": 12,
    },
    "homework": {
        "name": "课程作业",
        "description": "适用于日常课程作业、小论文",
        "paper_width": 21.0, "paper_height": 29.7,
        "top_margin": 2.54, "bottom_margin": 2.54,
        "left_margin": 3.17, "right_margin": 3.17,
        "body_font": "宋体", "body_en_font": "Times New Roman",
        "body_size": 12, "body_line_spacing": 1.5,
        "body_first_line_indent": 24, "body_alignment": "justify",
        "h1_font": "黑体", "h1_en_font": "Times New Roman",
        "h1_size": 14, "h1_bold": True, "h1_alignment": "left",
        "h1_space_before": 6, "h1_space_after": 6,
        "h2_font": "宋体", "h2_en_font": "Times New Roman",
        "h2_size": 12, "h2_bold": True, "h2_alignment": "left",
        "h2_space_before": 3, "h2_space_after": 3,
        "h3_font": "宋体", "h3_en_font": "Times New Roman",
        "h3_size": 12, "h3_bold": True, "h3_alignment": "left",
        "ref_font": "宋体", "ref_en_font": "Times New Roman",
        "ref_size": 10.5, "ref_line_spacing": 1.25,
        "keyword_font": "黑体", "keyword_en_font": "Times New Roman",
        "keyword_size": 12,
    },
    "journal": {
        "name": "期刊投稿格式",
        "description": "适用于期刊论文投稿",
        "paper_width": 21.0, "paper_height": 29.7,
        "top_margin": 2.54, "bottom_margin": 2.54,
        "left_margin": 2.54, "right_margin": 2.54,
        "body_font": "宋体", "body_en_font": "Times New Roman",
        "body_size": 10.5, "body_line_spacing": 1.25,
        "body_first_line_indent": 21, "body_alignment": "justify",
        "h1_font": "黑体", "h1_en_font": "Times New Roman",
        "h1_size": 14, "h1_bold": True, "h1_alignment": "center",
        "h1_space_before": 12, "h1_space_after": 6,
        "h2_font": "黑体", "h2_en_font": "Times New Roman",
        "h2_size": 12, "h2_bold": True, "h2_alignment": "left",
        "h2_space_before": 6, "h2_space_after": 3,
        "h3_font": "宋体", "h3_en_font": "Times New Roman",
        "h3_size": 10.5, "h3_bold": True, "h3_alignment": "left",
        "ref_font": "宋体", "ref_en_font": "Times New Roman",
        "ref_size": 9, "ref_line_spacing": 1.0,
        "keyword_font": "黑体", "keyword_en_font": "Times New Roman",
        "keyword_size": 10.5,
    },
}


def get_template_list():
    return [{"id": k, "name": v["name"], "description": v["description"]}
            for k, v in FORMAT_TEMPLATES.items()]


def get_template(template_id):
    return FORMAT_TEMPLATES.get(template_id, FORMAT_TEMPLATES["general"])


ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
ALIGN_CN = {
    WD_ALIGN_PARAGRAPH.LEFT: "左对齐",
    WD_ALIGN_PARAGRAPH.CENTER: "居中",
    WD_ALIGN_PARAGRAPH.RIGHT: "右对齐",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "两端对齐",
}


def _set_run_fonts(run, cn_font, en_font):
    """正确设置run的中英文字体"""
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)
    rFonts.set(qn('w:eastAsia'), cn_font)
    rFonts.set(qn('w:cs'), en_font)
    run.font.name = en_font


def _get_run_cn_font(run):
    """获取run的中文字体"""
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            return rFonts.get(qn('w:eastAsia'))
    return run.font.name


def _get_run_en_font(run):
    """获取run的英文字体"""
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            return rFonts.get(qn('w:ascii'))
    return run.font.name


class FormatIssue:
    def __init__(self, category, location, description, suggestion,
                 severity="warning", fixable=False):
        self.category = category
        self.location = location
        self.description = description
        self.suggestion = suggestion
        self.severity = severity
        self.fixable = fixable

    def to_dict(self):
        return {
            "category": self.category,
            "location": self.location,
            "description": self.description,
            "suggestion": self.suggestion,
            "severity": self.severity,
            "fixable": self.fixable,
        }


class FixRecord:
    """修复记录"""
    def __init__(self, category, location, item, old_value, new_value):
        self.category = category
        self.location = location
        self.item = item
        self.old_value = old_value
        self.new_value = new_value

    def to_dict(self):
        return {
            "category": self.category,
            "location": self.location,
            "item": self.item,
            "old_value": self.old_value,
            "new_value": self.new_value,
        }


class FormatChecker:
    def __init__(self, template_id="general", custom_config=None):
        if custom_config:
            self.config = custom_config
        else:
            self.config = get_template(template_id)
        self.template_id = template_id
        self.issues = []
        self.fix_records = []
        self.doc = None
        self.stats = {
            "total_paragraphs": 0, "total_words": 0,
            "h1_count": 0, "h2_count": 0, "h3_count": 0,
            "body_count": 0, "ref_count": 0, "table_count": 0,
        }
        self._para_types = []

    def check_file(self, file_path):
        self.issues = []
        self._para_types = []
        self.stats = {
            "total_paragraphs": 0, "total_words": 0,
            "h1_count": 0, "h2_count": 0, "h3_count": 0,
            "body_count": 0, "ref_count": 0, "table_count": 0,
        }
        self.doc = Document(file_path)
        self._check_page_setup()
        self._check_paragraphs()
        self._check_references()
        self._check_styles()
        self._check_tables()
        return self.issues, self.stats

    def fix_file(self, input_path, output_path):
        """自动修复，返回修复记录列表"""
        self.fix_records = []
        self.doc = Document(input_path)
        self._para_types = []

        for idx, para in enumerate(self.doc.paragraphs):
            self._para_types.append(self._get_paragraph_type(para))

        self._fix_page_setup()

        for idx, para in enumerate(self.doc.paragraphs):
            ptype = self._para_types[idx]
            loc = f"第{idx+1}段"
            if ptype == "h1":
                self._fix_heading(para, "h1", loc)
            elif ptype == "h2":
                self._fix_heading(para, "h2", loc)
            elif ptype == "h3":
                self._fix_heading(para, "h3", loc)
            elif ptype == "body":
                self._fix_body(para, loc)
            elif ptype == "reference":
                self._fix_reference(para, loc)
            elif ptype == "keyword":
                self._fix_keyword(para, loc)

        self._fix_tables()
        self.doc.save(output_path)
        return self.fix_records

    # ============================================================
    # 检查方法
    # ============================================================

    def _check_page_setup(self):
        c = self.config
        for i, section in enumerate(self.doc.sections):
            loc = f"第{i+1}节"
            if abs(section.page_width.cm - c["paper_width"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"纸张宽度为 {section.page_width.cm:.1f}cm，应为{c['paper_width']}cm",
                    "页面布局 → 纸张大小 → A4", "warning", True))
            if abs(section.page_height.cm - c["paper_height"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"纸张高度为 {section.page_height.cm:.1f}cm，应为{c['paper_height']}cm",
                    "页面布局 → 纸张大小 → A4", "warning", True))
            if abs(section.top_margin.cm - c["top_margin"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"上边距为 {section.top_margin.cm:.2f}cm，建议{c['top_margin']}cm",
                    "页面布局 → 页边距", "info", True))
            if abs(section.bottom_margin.cm - c["bottom_margin"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"下边距为 {section.bottom_margin.cm:.2f}cm，建议{c['bottom_margin']}cm",
                    "页面布局 → 页边距", "info", True))
            if abs(section.left_margin.cm - c["left_margin"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"左边距为 {section.left_margin.cm:.2f}cm，建议{c['left_margin']}cm",
                    "页面布局 → 页边距", "info", True))
            if abs(section.right_margin.cm - c["right_margin"]) > 0.1:
                self.issues.append(FormatIssue("页面设置", loc,
                    f"右边距为 {section.right_margin.cm:.2f}cm，建议{c['right_margin']}cm",
                    "页面布局 → 页边距", "info", True))

    def _get_paragraph_type(self, para):
        text = para.text.strip()
        if not text:
            return "empty"
        style_name = para.style.name if para.style else ""
        if re.match(r'^\[\d+\]', text):
            return "reference"
        if re.match(r'^[一二三四五六七八九十]+、', text) or re.match(r'^第[一二三四五六七八九十]+章', text):
            return "h1"
        if "Heading 1" in style_name or "标题 1" in style_name or "标题1" in style_name:
            return "h1"
        if re.match(r'^（[一二三四五六七八九十]+）', text) or re.match(r'^\d+\.\d+\s', text):
            return "h2"
        if "Heading 2" in style_name or "标题 2" in style_name or "标题2" in style_name:
            return "h2"
        if re.match(r'^\d+\.\s', text) or re.match(r'^\d+\.\d+\.\d+\s', text):
            return "h3"
        if "Heading 3" in style_name or "标题 3" in style_name or "标题3" in style_name:
            return "h3"
        if text.startswith("关键词") or text.startswith("摘 要") or text.startswith("摘要"):
            return "keyword"
        if text in ["参考文献", "References", "REFERENCES"]:
            return "ref_title"
        return "body"

    def _check_paragraphs(self):
        c = self.config
        for idx, para in enumerate(self.doc.paragraphs):
            text = para.text.strip()
            if not text:
                self._para_types.append("empty")
                continue
            self.stats["total_paragraphs"] += 1
            self.stats["total_words"] += len(text)
            ptype = self._get_paragraph_type(para)
            self._para_types.append(ptype)
            loc = f"第{idx+1}段"
            if ptype == "h1":
                self.stats["h1_count"] += 1
                self._check_heading(para, loc, "一级标题", "h1")
            elif ptype == "h2":
                self.stats["h2_count"] += 1
                self._check_heading(para, loc, "二级标题", "h2")
            elif ptype == "h3":
                self.stats["h3_count"] += 1
                self._check_heading(para, loc, "三级标题", "h3")
            elif ptype == "body":
                self.stats["body_count"] += 1
                self._check_body(para, loc, text)
            elif ptype == "reference":
                self.stats["ref_count"] += 1
                self._check_reference_item(para, loc, text)
            elif ptype == "keyword":
                self._check_keyword(para, loc, text)

    def _check_heading(self, para, loc, title, level_key):
        c = self.config
        pf = para.paragraph_format
        expected_align = ALIGN_MAP[c[f"{level_key}_alignment"]]

        if pf.alignment != expected_align:
            actual = ALIGN_CN.get(pf.alignment, str(pf.alignment))
            expected = ALIGN_CN.get(expected_align, str(expected_align))
            self.issues.append(FormatIssue(f"{title}格式", loc,
                f"对齐方式为{actual}，应为{expected}",
                f"设置对齐方式为{expected}", "warning", True))

        for run in para.runs:
            if not run.text.strip():
                continue
            cn_font = _get_run_cn_font(run)
            en_font = _get_run_en_font(run)
            expected_cn = c[f"{level_key}_font"]
            expected_en = c[f"{level_key}_en_font"]

            if cn_font and cn_font != expected_cn:
                self.issues.append(FormatIssue(f"{title}格式", loc,
                    f"中文字体为「{cn_font}」，应为「{expected_cn}」",
                    f"中文字体改为{expected_cn}", "warning", True))
            if en_font and en_font != expected_en:
                self.issues.append(FormatIssue(f"{title}格式", loc,
                    f"英文/数字字体为「{en_font}」，应为「{expected_en}」",
                    f"英文字体改为{expected_en}", "info", True))
            if run.font.size and abs(run.font.size.pt - c[f"{level_key}_size"]) > 0.5:
                self.issues.append(FormatIssue(f"{title}格式", loc,
                    f"字号为{run.font.size.pt:.1f}pt，应为{c[f'{level_key}_size']:.1f}pt",
                    f"字号改为{c[f'{level_key}_size']:.1f}pt", "warning", True))
            if run.font.bold != c[f"{level_key}_bold"]:
                status = "加粗" if run.font.bold else "未加粗"
                should = "加粗" if c[f"{level_key}_bold"] else "不加粗"
                self.issues.append(FormatIssue(f"{title}格式", loc,
                    f"标题{status}，应{should}",
                    f"设置{'加粗' if c[f'{level_key}_bold'] else '取消加粗'}", "info", True))

        sb = c.get(f"{level_key}_space_before")
        sa = c.get(f"{level_key}_space_after")
        if sb is not None and pf.space_before and abs(pf.space_before.pt - sb) > 1:
            self.issues.append(FormatIssue(f"{title}格式", loc,
                f"段前间距为{pf.space_before.pt:.1f}pt，建议{sb}pt",
                f"段前设为{sb}pt", "info", True))
        if sa is not None and pf.space_after and abs(pf.space_after.pt - sa) > 1:
            self.issues.append(FormatIssue(f"{title}格式", loc,
                f"段后间距为{pf.space_after.pt:.1f}pt，建议{sa}pt",
                f"段后设为{sa}pt", "info", True))

    def _check_body(self, para, loc, text):
        c = self.config
        pf = para.paragraph_format

        if pf.first_line_indent:
            if abs(pf.first_line_indent.pt - c["body_first_line_indent"]) > 2:
                self.issues.append(FormatIssue("正文格式", loc,
                    f"首行缩进为{pf.first_line_indent.pt:.1f}pt，应为{c['body_first_line_indent']}pt（2字符）",
                    "首行缩进设为2字符", "warning", True))
        elif len(text) > 20:
            self.issues.append(FormatIssue("正文格式", loc,
                "未设置首行缩进", "首行缩进设为2字符", "warning", True))

        if pf.line_spacing:
            if isinstance(pf.line_spacing, float):
                if abs(pf.line_spacing - c["body_line_spacing"]) > 0.1:
                    self.issues.append(FormatIssue("正文格式", loc,
                        f"行距为{pf.line_spacing:.1f}倍，建议{c['body_line_spacing']}倍",
                        f"行距设为{c['body_line_spacing']}倍", "warning", True))
        else:
            self.issues.append(FormatIssue("正文格式", loc,
                "未设置明确行距", f"行距设为{c['body_line_spacing']}倍", "info", True))

        expected_align = ALIGN_MAP[c["body_alignment"]]
        if pf.alignment != expected_align:
            actual = ALIGN_CN.get(pf.alignment, str(pf.alignment))
            self.issues.append(FormatIssue("正文格式", loc,
                f"对齐方式为{actual}，建议两端对齐",
                "对齐方式设为两端对齐", "info", True))

        for run in para.runs:
            if not run.text.strip():
                continue
            cn_font = _get_run_cn_font(run)
            en_font = _get_run_en_font(run)
            has_cn = bool(re.search(r'[\u4e00-\u9fff]', run.text))
            has_en = bool(re.search(r'[a-zA-Z0-9]', run.text))

            if has_cn and cn_font and cn_font != c["body_font"]:
                self.issues.append(FormatIssue("正文字体", loc,
                    f"中文字体为「{cn_font}」，应为「{c['body_font']}」",
                    f"中文字体改为{c['body_font']}", "warning", True))
            if has_en and en_font and en_font != c["body_en_font"]:
                self.issues.append(FormatIssue("正文字体", loc,
                    f"英文/数字字体为「{en_font}」，应为「{c['body_en_font']}」",
                    f"英文字体改为{c['body_en_font']}", "info", True))
            if run.font.size and abs(run.font.size.pt - c["body_size"]) > 0.5:
                self.issues.append(FormatIssue("正文字号", loc,
                    f"字号为{run.font.size.pt:.1f}pt，建议{c['body_size']:.1f}pt",
                    f"字号改为{c['body_size']:.1f}pt", "warning", True))

    def _check_reference_item(self, para, loc, text):
        c = self.config
        for run in para.runs:
            if not run.text.strip():
                continue
            cn_font = _get_run_cn_font(run)
            if cn_font and cn_font != c["ref_font"]:
                self.issues.append(FormatIssue("参考文献格式", loc,
                    f"中文字体为「{cn_font}」，建议「{c['ref_font']}」",
                    f"字体改为{c['ref_font']}", "info", True))
            if run.font.size and abs(run.font.size.pt - c["ref_size"]) > 0.5:
                self.issues.append(FormatIssue("参考文献格式", loc,
                    f"字号为{run.font.size.pt:.1f}pt，建议{c['ref_size']:.1f}pt",
                    f"字号改为{c['ref_size']:.1f}pt", "info", True))
        if not re.match(r'^\[\d+\]\s*.+\..+\.', text):
            self.issues.append(FormatIssue("参考文献格式", loc,
                "格式可能不规范，标准格式：[序号] 作者. 题名[J]. 刊名, 年, 卷(期): 页码.",
                "参照GB/T 7714标准调整", "warning", False))

    def _check_keyword(self, para, loc, text):
        c = self.config
        for run in para.runs:
            if not run.text.strip():
                continue
            if "关键词" in run.text or "摘要" in run.text:
                cn_font = _get_run_cn_font(run)
                if cn_font and cn_font != c["keyword_font"]:
                    self.issues.append(FormatIssue("摘要关键词", loc,
                        f"标签字体为「{cn_font}」，建议「{c['keyword_font']}」",
                        f"字体改为{c['keyword_font']}", "info", True))

    def _check_references(self):
        ref_items = []
        for para in self.doc.paragraphs:
            text = para.text.strip()
            if re.match(r'^\[\d+\]', text):
                ref_items.append(text)
        if not ref_items:
            self.issues.append(FormatIssue("参考文献", "全文",
                "未检测到参考文献条目", "如需要请添加参考文献", "info", False))
            return
        numbers = []
        for text in ref_items:
            m = re.match(r'^\[(\d+)\]', text)
            if m:
                numbers.append(int(m.group(1)))
        if numbers and numbers != list(range(1, len(numbers) + 1)):
            self.issues.append(FormatIssue("参考文献", "全文",
                f"参考文献序号不连续，当前序号：{numbers}",
                "将序号调整为从[1]开始的连续编号", "warning", False))

    def _check_styles(self):
        heading_used = any(
            para.style and "Heading" in para.style.name
            for para in self.doc.paragraphs
        )
        if not heading_used and self.stats["h1_count"] > 0:
            self.issues.append(FormatIssue("样式规范", "全文",
                "标题未使用Word内置标题样式",
                "建议使用「开始 → 样式」中的标题1/2/3", "info", False))

    def _check_tables(self):
        self.stats["table_count"] = len(self.doc.tables)
        for i, table in enumerate(self.doc.tables):
            if table.rows:
                for cell in table.rows[0].cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text.strip() and not run.font.bold:
                                self.issues.append(FormatIssue("表格格式", f"第{i+1}个表格",
                                    "表头单元格未加粗", "选中表头行 → 设置加粗", "info", True))
                                return

    # ============================================================
    # 修复方法（带修改记录）
    # ============================================================

    def _fix_page_setup(self):
        c = self.config
        for i, section in enumerate(self.doc.sections):
            loc = f"第{i+1}节"
            old_w = section.page_width.cm
            section.page_width = Cm(c["paper_width"])
            if abs(old_w - c["paper_width"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "纸张宽度", f"{old_w:.1f}cm", f"{c['paper_width']}cm"))
            old_h = section.page_height.cm
            section.page_height = Cm(c["paper_height"])
            if abs(old_h - c["paper_height"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "纸张高度", f"{old_h:.1f}cm", f"{c['paper_height']}cm"))
            old_t = section.top_margin.cm
            section.top_margin = Cm(c["top_margin"])
            if abs(old_t - c["top_margin"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "上边距", f"{old_t:.2f}cm", f"{c['top_margin']}cm"))
            old_b = section.bottom_margin.cm
            section.bottom_margin = Cm(c["bottom_margin"])
            if abs(old_b - c["bottom_margin"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "下边距", f"{old_b:.2f}cm", f"{c['bottom_margin']}cm"))
            old_l = section.left_margin.cm
            section.left_margin = Cm(c["left_margin"])
            if abs(old_l - c["left_margin"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "左边距", f"{old_l:.2f}cm", f"{c['left_margin']}cm"))
            old_r = section.right_margin.cm
            section.right_margin = Cm(c["right_margin"])
            if abs(old_r - c["right_margin"]) > 0.1:
                self.fix_records.append(FixRecord("页面设置", loc, "右边距", f"{old_r:.2f}cm", f"{c['right_margin']}cm"))

    def _fix_heading(self, para, level_key, loc):
        c = self.config
        pf = para.paragraph_format
        title_cn = {"h1": "一级标题", "h2": "二级标题", "h3": "三级标题"}[level_key]

        old_align = ALIGN_CN.get(pf.alignment, "未设置")
        pf.alignment = ALIGN_MAP[c[f"{level_key}_alignment"]]
        if old_align != ALIGN_CN[ALIGN_MAP[c[f"{level_key}_alignment"]]]:
            self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "对齐方式", old_align, ALIGN_CN[ALIGN_MAP[c[f"{level_key}_alignment"]]]))

        sb = c.get(f"{level_key}_space_before")
        if sb is not None:
            old_sb = pf.space_before.pt if pf.space_before else 0
            pf.space_before = Pt(sb)
            if abs(old_sb - sb) > 1:
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "段前间距", f"{old_sb:.1f}pt", f"{sb}pt"))
        sa = c.get(f"{level_key}_space_after")
        if sa is not None:
            old_sa = pf.space_after.pt if pf.space_after else 0
            pf.space_after = Pt(sa)
            if abs(old_sa - sa) > 1:
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "段后间距", f"{old_sa:.1f}pt", f"{sa}pt"))

        pf.first_line_indent = Pt(0)

        for run in para.runs:
            if not run.text.strip():
                continue
            old_cn = _get_run_cn_font(run) or "未设置"
            old_en = _get_run_en_font(run) or "未设置"
            old_size = f"{run.font.size.pt:.1f}pt" if run.font.size else "未设置"
            old_bold = "加粗" if run.font.bold else "不加粗"

            _set_run_fonts(run, c[f"{level_key}_font"], c[f"{level_key}_en_font"])
            run.font.size = Pt(c[f"{level_key}_size"])
            run.font.bold = c[f"{level_key}_bold"]

            if old_cn != c[f"{level_key}_font"]:
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "中文字体", old_cn, c[f"{level_key}_font"]))
            if old_en != c[f"{level_key}_en_font"]:
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "英文/数字字体", old_en, c[f"{level_key}_en_font"]))
            if old_size != f"{c[f'{level_key}_size']:.1f}pt":
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "字号", old_size, f"{c[f'{level_key}_size']:.1f}pt"))
            new_bold = "加粗" if c[f"{level_key}_bold"] else "不加粗"
            if old_bold != new_bold:
                self.fix_records.append(FixRecord(f"{title_cn}格式", loc, "加粗", old_bold, new_bold))

    def _fix_body(self, para, loc):
        c = self.config
        pf = para.paragraph_format

        old_indent = f"{pf.first_line_indent.pt:.1f}pt" if pf.first_line_indent else "未设置"
        pf.first_line_indent = Pt(c["body_first_line_indent"])
        if old_indent != f"{c['body_first_line_indent']:.1f}pt":
            self.fix_records.append(FixRecord("正文格式", loc, "首行缩进", old_indent, f"{c['body_first_line_indent']}pt"))

        old_line = f"{pf.line_spacing:.1f}倍" if isinstance(pf.line_spacing, float) else "未设置"
        pf.line_spacing = c["body_line_spacing"]
        if old_line != f"{c['body_line_spacing']:.1f}倍":
            self.fix_records.append(FixRecord("正文格式", loc, "行距", old_line, f"{c['body_line_spacing']}倍"))

        old_align = ALIGN_CN.get(pf.alignment, "未设置")
        pf.alignment = ALIGN_MAP[c["body_alignment"]]
        if old_align != ALIGN_CN[ALIGN_MAP[c["body_alignment"]]]:
            self.fix_records.append(FixRecord("正文格式", loc, "对齐方式", old_align, ALIGN_CN[ALIGN_MAP[c["body_alignment"]]]))

        for run in para.runs:
            if not run.text.strip():
                continue
            old_cn = _get_run_cn_font(run) or "未设置"
            old_en = _get_run_en_font(run) or "未设置"
            old_size = f"{run.font.size.pt:.1f}pt" if run.font.size else "未设置"

            _set_run_fonts(run, c["body_font"], c["body_en_font"])
            run.font.size = Pt(c["body_size"])

            has_cn = bool(re.search(r'[\u4e00-\u9fff]', run.text))
            has_en = bool(re.search(r'[a-zA-Z0-9]', run.text))
            if has_cn and old_cn != c["body_font"]:
                self.fix_records.append(FixRecord("正文字体", loc, "中文字体", old_cn, c["body_font"]))
            if has_en and old_en != c["body_en_font"]:
                self.fix_records.append(FixRecord("正文字体", loc, "英文/数字字体", old_en, c["body_en_font"]))
            if old_size != f"{c['body_size']:.1f}pt":
                self.fix_records.append(FixRecord("正文字号", loc, "字号", old_size, f"{c['body_size']:.1f}pt"))

    def _fix_reference(self, para, loc):
        c = self.config
        pf = para.paragraph_format
        pf.line_spacing = c["ref_line_spacing"]
        pf.first_line_indent = Pt(0)
        for run in para.runs:
            old_cn = _get_run_cn_font(run) or "未设置"
            old_size = f"{run.font.size.pt:.1f}pt" if run.font.size else "未设置"
            _set_run_fonts(run, c["ref_font"], c["ref_en_font"])
            run.font.size = Pt(c["ref_size"])
            if old_cn != c["ref_font"]:
                self.fix_records.append(FixRecord("参考文献格式", loc, "字体", old_cn, c["ref_font"]))
            if old_size != f"{c['ref_size']:.1f}pt":
                self.fix_records.append(FixRecord("参考文献格式", loc, "字号", old_size, f"{c['ref_size']:.1f}pt"))

    def _fix_keyword(self, para, loc):
        c = self.config
        for run in para.runs:
            if "关键词" in run.text or "摘要" in run.text:
                old_cn = _get_run_cn_font(run) or "未设置"
                _set_run_fonts(run, c["keyword_font"], c["keyword_en_font"])
                run.font.size = Pt(c["keyword_size"])
                run.font.bold = True
                if old_cn != c["keyword_font"]:
                    self.fix_records.append(FixRecord("摘要关键词", loc, "标签字体", old_cn, c["keyword_font"]))

    def _fix_tables(self):
        for i, table in enumerate(self.doc.tables):
            if table.rows:
                for cell in table.rows[0].cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text.strip() and not run.font.bold:
                                run.font.bold = True
                                self.fix_records.append(FixRecord("表格格式", f"第{i+1}个表格", "表头加粗", "未加粗", "加粗"))

    def get_summary(self):
        return {
            "total_issues": len(self.issues),
            "error_count": sum(1 for i in self.issues if i.severity == "error"),
            "warning_count": sum(1 for i in self.issues if i.severity == "warning"),
            "info_count": sum(1 for i in self.issues if i.severity == "info"),
            "fixable_count": sum(1 for i in self.issues if i.fixable),
            "category_stats": self._count_by_category(),
            "stats": self.stats,
        }

    def _count_by_category(self):
        result = {}
        for issue in self.issues:
            result[issue.category] = result.get(issue.category, 0) + 1
        return result
